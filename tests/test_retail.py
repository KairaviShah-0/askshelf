"""Tests for the upload-driven engine: ingestion, privacy, safe SQL, column mapping, every analysis, alerts, forecasting and the report writers."""
import io
from pathlib import Path

import pandas as pd
import pytest

from agent import datalab as D
from agent import labai, labreport, rbac, report_docx, report_pdf, report_xlsx
from agent import retail as R

SAMPLE = Path(__file__).resolve().parent.parent / "sample_data" / "retail_orders_sample.csv"


@pytest.fixture(scope="module")
def lab():
    lb = D.Lab("pytest")
    lb.add_file(SAMPLE.name, SAMPLE.read_bytes())
    return lb


def mini_lab(df: pd.DataFrame, name="x.csv") -> D.Lab:
    lb = D.Lab("mini")
    lb.add_file(name, df.to_csv(index=False).encode())
    return lb


# ---------------------------------------------------------------------------- ingestion and privacy
def test_sample_loads_and_personal_column_is_removed(lab):
    t = next(iter(lab.tables.values()))
    assert len(t) > 4000 and "customer_email" not in t.columns and [d["column"] for d in lab.dropped] == ["customer_email"]


def test_pii_detected_by_values_not_just_names():
    lb = mini_lab(pd.DataFrame({"contact": [f"user{i}@example.com" for i in range(20)], "amount": range(20)}))
    assert "contact" not in next(iter(lb.tables.values())).columns


def test_excel_and_messy_values_are_cleaned():
    buf = io.BytesIO()
    pd.DataFrame({"Order Date": ["2026-01-05", "2026-02-03"] * 5, "Sales $": ["$1,200", "$900"] * 5, "Region": ["N", "S"] * 5}).to_excel(buf, index=False)
    lb = D.Lab("xl")
    lb.add_file("Book 1.xlsx", buf.getvalue())
    t = next(iter(lb.tables.values()))
    assert list(t.columns) == ["order_date", "sales", "region"] and t["sales"].sum() == 10500 and str(t["order_date"].dtype).startswith("datetime")


def test_files_with_same_columns_are_fused():
    df = pd.DataFrame({"order_date": ["2026-01-01"] * 6, "sales": range(6)})
    lb = D.Lab("fuse")
    lb.add_file("a.csv", df.to_csv(index=False).encode())
    lb.add_file("b.csv", df.to_csv(index=False).encode())
    assert len(lb.tables) == 1 and len(next(iter(lb.tables.values()))) == 12 and "combined" in lb.fusion_note


@pytest.mark.parametrize("sql", ["DROP TABLE retail_orders_sample", "SELECT 1; DELETE FROM retail_orders_sample", "PRAGMA table_info(x)", "SELECT * FROM sqlite_master", "UPDATE retail_orders_sample SET store='x'", "SELECT load_extension('x')"])
def test_unsafe_sql_is_refused(lab, sql):
    with pytest.raises((D.UnsafeSQL, ValueError)):
        D.run_select(lab, sql)
    assert len(lab.tables["retail_orders_sample"]) > 4000


# ---------------------------------------------------------------------------- mapping
def test_mapping_is_detected_and_editable(lab):
    m = R.get_mapping(lab)["map"]
    assert m["date"] == "order_date" and m["amount"] == "sales_amount" and m["due"] == "due_date" and m["delivered"] == "delivered_date" and m["stock"] == "stock_on_hand"
    assert R.capabilities(lab)["stock"] and R.capabilities(lab)["margin"]


def test_missing_columns_switch_pages_off():
    lb = mini_lab(pd.DataFrame({"order_date": pd.date_range("2026-01-01", periods=40).strftime("%Y-%m-%d"), "sales": range(40), "store": ["A", "B"] * 20}))
    caps = R.capabilities(lb)
    assert caps["dashboard"] and caps["forecasts"] and not caps["deliveries"] and not caps["stock"] and not caps["penalties"]
    with pytest.raises(ValueError):
        R.set_mapping(lb, "x", {"date": "store"})
    lab2 = D.Lab("m2")
    lab2.add_file("x.csv", pd.DataFrame({"order_date": ["2026-01-01"] * 8, "sales": range(8)}).to_csv(index=False).encode())
    R.set_mapping(lab2, "x", {"date": "order_date", "amount": None})
    assert not R.capabilities(lab2)["dashboard"]


# ---------------------------------------------------------------------------- analyses
def test_dashboard(lab):
    d = R.dashboard(lab, "90d")
    labels = [k["label"] for k in d["kpis"]]
    assert {"Revenue", "Orders", "Gross margin", "On-time delivery"} <= set(labels) and len(d["charts"]) >= 4 and d["compared"] and d["insights"]
    assert R.dashboard(lab, "all")["compared"] is False


def test_deliveries_model_is_trained_on_uploaded_history(lab):
    r = R.deliveries(lab)
    assert r["model_ok"] and 0.5 < r["model"]["auc"] <= 1 and r["open_table"]["total_rows"] > 0 and r["n_open"] > 0
    small = mini_lab(pd.DataFrame({"order_date": ["2026-01-01"] * 30, "sales": range(30), "due_date": ["2026-01-05"] * 30, "delivered_date": ["2026-01-08"] * 30}))
    r2 = R.deliveries(small)
    assert not r2["model_ok"] and any("at least" in n for n in r2["notes"])


def test_stock(lab):
    r = R.stock(lab, 21)
    assert any(k["label"] == "Out of stock" for k in r["kpis"]) and r["table"]["total_rows"] == 30 and "Suggested order" in r["table"]["columns"]


def test_forecast_and_groups(lab):
    r = R.forecast(lab, "units", "category", 4)
    assert r["unit"] == "month" and r["group_chart"] and len(r["table"]["data"]) == 4 and r["metric"] == "units"


def test_planner_math(lab):
    base = R.planner(lab, 0, 0, 0, "none")
    up = R.planner(lab, 10, 0, 0, "none")
    rev = lambda r: float(next(k for k in r["kpis"] if k["label"] == "Monthly revenue")["value"].replace(",", "").replace("K", "e3").replace("M", "e6"))
    assert rev(up) > rev(base) and base["has_cost"] and R.planner(lab, 10, 0, 0, "high")["kpis"][0]["delta"] is not None


def test_penalties_plan_saves_money_and_assumed_rate(lab):
    p0, p5 = R.penalties(lab, 0), R.penalties(lab, 5)
    val = lambda r, label: next((k for k in r["kpis"] if k["label"] == label), None)
    assert val(p0, "After your plan") is None and val(p5, "After your plan") is not None and p5["plan_ok"] and p5["source"] == "mapped"
    df = pd.read_csv(SAMPLE).drop(columns=["late_penalty", "customer_email"])
    lb = mini_lab(df)
    assert R.penalties(lb, 0, 0)["source"] == "none" and R.penalties(lb, 0, 5)["source"] == "assumed"


def test_alerts_are_actionable_and_adapt_to_columns(lab):
    al = R.alerts(lab)
    assert al and all(a["action"] and a["ask"] and a["page"] in ("sales", "delivery", "stock") for a in al)
    cats = {a["category"] for a in al}
    assert {"Stock-out / low stock", "Overdue orders"} <= cats
    lb = mini_lab(pd.DataFrame({"order_date": pd.date_range("2025-01-01", periods=300).strftime("%Y-%m-%d"), "sales": [100] * 270 + [10] * 30}))
    assert not {"Overdue orders", "Stock-out / low stock"} & {a["category"] for a in R.alerts(lb)}


# ---------------------------------------------------------------------------- questions, roles, reports
def test_rule_based_questions(lab):
    top = labai.lab_answer(lab, "Top 3 suppliers by sales amount", [], None)
    assert top["kind"] == "sql" and top["table"]["total_rows"] == 3 and top["chart"]
    assert labai.lab_answer(lab, "show me customer ID 45", [], None)["blocked"]
    assert "Upload" in labai.lab_answer(D.Lab("empty"), "top suppliers", [], None)["error"]


def test_roles_gate_pages():
    assert rbac.can_view("Operations Manager", "stock") and not rbac.can_view("Operations Manager", "dashboard") and not rbac.can_view("Sales Manager", "reports")
    assert rbac.ROLES["Operations Manager"].can_upload and not rbac.ROLES["Sales Manager"].can_upload and not rbac.ROLES["Operations Manager"].can_export
    assert len(rbac.access_matrix()) == len(rbac.PAGES) + 2


@pytest.fixture(scope="module")
def report(lab):
    return labreport.build_report(lab, None, {"organisation": "Test Co", "prepared_for": "Board"})


def test_report_uses_mapped_columns_and_no_sample_company_names(report):
    assert len(report.sections) >= 5 and report.forecast_table is not None and report.recommendations and report.references
    text = " ".join(report.exec_summary + report.conclusions + [s.title for s in report.sections])
    assert "sales amount" in text.lower() and "furniture" not in text.lower()


def test_report_pdf_docx_xlsx(report):
    pdf, pages = report_pdf.build_pdf(report)
    assert pdf[:4] == b"%PDF" and len(pdf) > 100_000 and any("Executive" in k for k in pages)
    from docx import Document
    d = Document(io.BytesIO(report_docx.build_docx(report, pages)))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "Executive Summary" in heads and "References" in heads and len(d.inline_shapes) == len(report.figures)
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(report_xlsx.report_xlsx(report)))
    assert wb.sheetnames[0] == "Summary" and "Queries" in wb.sheetnames


def test_no_furniture_left_in_the_product():
    root = Path(__file__).resolve().parent.parent
    banned = ("furniture", "timber", "walnut", "modern furniture")
    for f in list((root / "agent").glob("*.py")) + [root / "server.py"] + list((root / "web").rglob("*.*")) + list((root / "sample_data").glob("*.csv")):
        if f.suffix in (".py", ".html", ".js", ".css", ".csv"):
            assert not any(b in f.read_text(errors="ignore").lower() for b in banned), f
